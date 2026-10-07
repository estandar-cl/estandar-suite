from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "plugins/estandar-leads/docs/ESTANDAR_GOOGLE_ADS_API_TOOL_DESIGN.docx"
BLUE = RGBColor(46, 116, 181)
GRAY = RGBColor(90, 90, 90)


def set_font(run, size=10.5, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6 if level == 1 else 4)
    r = p.add_run(text)
    set_font(r, 15 if level == 1 else 12.5, True, BLUE)
    return p


def add_body(doc, text, bold_label=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    if bold_label and text.startswith(bold_label):
        r = p.add_run(bold_label)
        set_font(r, 10.5, True)
        r = p.add_run(text[len(bold_label):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.10
        set_font(p.add_run(item))


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = Inches(widths[idx])
        set_cell_shading(cell, "F2F4F7")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_font(cell.paragraphs[0].add_run(header), 9.5, True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].width = Inches(widths[idx])
            set_cell_margins(cells[idx])
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_font(cells[idx].paragraphs[0].add_run(value), 9.5)
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(header.add_run("Estandar Studio | Google Ads API"), 9, False, GRAY)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run("Confidential technical design | June 11, 2026"), 9, False, GRAY)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(20)
p.paragraph_format.space_after = Pt(4)
set_font(p.add_run("GOOGLE ADS API TOOL DESIGN"), 23, True)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(18)
set_font(p.add_run("Estandar Leads Offline Conversion Import"), 14, False, GRAY)

for label, value in (
    ("Company", "Estandar Studio"),
    ("Website", "https://estandar.cl"),
    ("API contact", "gestion@estandar.cl"),
    ("Manager account", "979-364-5349"),
    ("Tool users", "Internal Estandar Studio operators only"),
):
    add_body(doc, f"{label}: {value}", f"{label}:")

add_heading(doc, "1. Purpose")
add_body(
    doc,
    "The tool connects WordPress lead-management installations operated by Estandar Studio "
    "to the Google Ads API. Its sole API purpose is to upload verified offline conversions "
    "for authorized client accounts after a real commercial milestone occurs."
)
add_body(
    doc,
    "The tool does not create campaigns, alter bids, scrape Google Ads, expose API access to "
    "clients, or provide a public self-service interface."
)

add_heading(doc, "2. Business workflow")
add_bullets(doc, [
    "A visitor arrives through an authorized client's Google Ads campaign.",
    "The site stores the click identifier and attribution data with the lead record.",
    "An authorized operator validates the lead and updates its commercial status.",
    "Only configured qualifying statuses create an idempotent offline-conversion queue item.",
    "A scheduled worker uploads eligible click conversions to the client's conversion action.",
    "The queue records success, retryable failure, or dead-letter status for operational review.",
])

doc.add_page_break()
add_heading(doc, "3. Architecture and data flow")
add_table(doc, ["Component", "Responsibility"], [
    ("Client website", "Captures consented first-party lead and attribution data."),
    ("Estandar Leads", "Maintains canonical lead status and validates qualifying milestones."),
    ("OCI queue", "Persists idempotent jobs and retry state in the client WordPress database."),
    ("OAuth 2.0", "Obtains short-lived access tokens from a securely stored refresh token."),
    ("Google Ads API", "Receives click conversion uploads for the authorized customer account."),
], [1.55, 4.75])

add_heading(doc, "4. API operation")
add_body(doc, "The integration uses the Google Ads API click-conversion upload service. Each request is scoped to a configured client customer ID and conversion action.")
add_table(doc, ["Field", "Source / rule"], [
    ("gclid", "First-party attribution captured from the client's landing session."),
    ("conversion_action", "Configured Google Ads conversion action resource."),
    ("conversion_date_time", "UTC commercial status-change time with timezone offset."),
    ("conversion_value", "Optional configured value; no personal data is transmitted."),
    ("currency_code", "Configured ISO currency when a conversion value is used."),
], [1.75, 4.55])

add_heading(doc, "5. Qualification rules")
add_body(doc, "A form submission alone is retained as a secondary measurement signal. The primary optimization signal is generated only after human or business-system validation.")
add_table(doc, ["Lead status", "Upload behavior"], [
    ("NEW / CONTACTED", "Not uploaded"),
    ("QUALIFIED", "Uploaded as a validated lead"),
    ("MEETING_BOOKED", "Uploaded as validated commercial intent"),
    ("CONVERTED", "Uploaded as a won customer"),
    ("Invalid, junk, spam, education", "Not uploaded"),
], [2.05, 4.25])

add_heading(doc, "6. Authentication and access control")
add_bullets(doc, [
    "Google OAuth 2.0 is used with a client ID, client secret, and refresh token.",
    "Credentials are stored in protected WordPress options and are never returned to public pages.",
    "Only authenticated WordPress administrators can view or change integration settings.",
    "The Google Ads developer token belongs to the manager account; client customer IDs are explicitly configured.",
    "The tool is used only by Estandar Studio employees or authorized contractors.",
])

add_heading(doc, "7. Reliability controls")
add_bullets(doc, [
    "Unique lead/status keys prevent duplicate queue creation.",
    "The worker processes a bounded batch of up to 20 records per scheduled run.",
    "Failed uploads retry at most three times with a minimum one-hour delay.",
    "Exhausted records remain in a dead-letter state for inspection.",
    "Test mode logs payload preparation without calling the Google Ads API.",
])

add_heading(doc, "8. Privacy and policy controls")
add_bullets(doc, [
    "Only click identifiers and conversion metadata required by Google Ads are uploaded.",
    "Names, email addresses, phone numbers, messages, and tax/accounting information are not sent.",
    "Uploads occur only for client accounts contractually authorized to use the integration.",
    "Client-site consent and privacy notices govern first-party attribution capture.",
    "Data is retained within each client's WordPress database according to the client's retention policy.",
])

add_heading(doc, "9. Deployment scope")
add_body(
    doc,
    "Initial deployment is for AK Contadores and subsequent Estandar Studio clients that explicitly "
    "authorize the integration. Every installation has separate credentials, customer ID, conversion "
    "action, queue, and audit history. There is no cross-client lead database."
)

add_heading(doc, "10. Support")
add_body(doc, "Operational and API notices: gestion@estandar.cl")

doc.save(OUT)
